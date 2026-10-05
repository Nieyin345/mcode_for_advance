import sys, pathlib, zipfile
from pptx import Presentation
from openpyxl import Workbook
root=pathlib.Path(sys.argv[1])
p=Presentation()
for i in range(2):
 slide=p.slides.add_slide(p.slide_layouts[5])
 slide.shapes.title.text="SLIDE_%d 研究😀 "%(i+1)+"Research text "*100
p.save(root/"slides.pptx")
w=Workbook();w.active.title="Research";w.active.append(["Title","Year"]);w.active.append(["研究论文😀",2026]);w.save(root/"table.xlsx")
xml='<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>'+('研究😀 文档正文 '*300)+'</w:t></w:r></w:p></w:body></w:document>'
with zipfile.ZipFile(root/"text.docx","w") as z:z.writestr("word/document.xml",xml)
print("Generated isolated PPTX/XLSX/DOCX fixtures")
