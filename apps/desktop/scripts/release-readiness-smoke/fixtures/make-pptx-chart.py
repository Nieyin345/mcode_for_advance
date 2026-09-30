"""Optional fixture regeneration: pip install python-pptx==1.0.2.

The smoke only reads the committed JSON, so running tests does not need Python,
Office, a model, a user document or network access.
"""
import base64
import json
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = "Release chart fixture"
data = CategoryChartData()
data.categories = ["Alpha", "Beta"]
data.add_series("Series A", (10, 20))
slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(1), Inches(1.5), Inches(8), Inches(4.5), data,
)
buffer = BytesIO()
presentation.save(buffer)
Path(__file__).with_name("pptx-chart.json").write_text(
    json.dumps({
        "description": "Synthetic one-slide PPTX created with python-pptx 1.0.2; title, two categories and one bar series. No user content or external relationships.",
        "base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
    }, indent=2) + "\n",
    encoding="utf-8",
)
