import sys, pathlib, zipfile
from openpyxl import Workbook

root = pathlib.Path(sys.argv[1])

# 一个"够大"的工作簿:读取越慢,并发读-改-写的窗口越宽(MC-069 的 E1 夹具形状)。
# A1/B1 是两个**互不重叠**的目标格,用来证明并发编辑会各写一份完整快照、互相覆盖。
w = Workbook()
ws = w.active
ws.title = "Audit"
ws["A1"] = "ORIGINAL_A"
ws["B1"] = "ORIGINAL_B"
for r in range(2, 2001):
    for c in range(1, 9):
        ws.cell(r, c).value = r * c
w.save(root / "audit.xlsx")

# 一个 DOCX:正文里有两个互不重叠的标记,供并发 XML 编辑。
xml = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:body><w:p><w:r><w:t>MARK_A and MARK_B</w:t></w:r></w:p></w:body></w:document>'
)
with zipfile.ZipFile(root / "doc.docx", "w") as z:
    z.writestr("word/document.xml", xml)

print("doc-ops fixtures ready")
