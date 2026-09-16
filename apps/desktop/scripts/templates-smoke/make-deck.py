# -*- coding: utf-8 -*-
"""造一份**真的**演示稿,用来核对 pptx 在应用内排成什么样。

    python scripts/templates-smoke/make-deck.py            # → deck.pptx(基准)
    python scripts/templates-smoke/make-deck.py --hard     # → 再加图片 / 图表 / 渐变

## 它不属于冒烟套件

`make-fixtures.py`(只用标准库)造的那三个包是给 `run.sh` 跑的 —— 那边只验"字节对不对
路",不验渲染。这个脚本是**另一件事**:给你一份我能想到的、最像真模版的稿子,拿去在真
Chromium 里开一眼。所以它可以用第三方库(Pillow / python-pptx),不进 `run.sh`。

为什么要用 python-pptx 而不是手拼 ZIP:`pptx-preview` 要读母版 / 版式 / 主题
(`ppt/slideMasters`、`ppt/slideLayouts`、`ppt/theme`),手拼一个"只有一个 slideN.xml"
的包它排不出来 —— 那不是库的问题,是夹具不像真的。python-pptx 写出来的就是 PowerPoint
自己认的包。

## 怎么"在真 Chromium 里开一眼"

渲染端这一层没有无头测试(要真 DOM)。但**不用装 Playwright**,三步就够:

1. 把库的产物拷进一个临时目录 —— `node_modules/.pnpm/pptx-preview@*/node_modules/
   pptx-preview/dist/pptx-preview.umd.js`(UMD,浏览器直接 `<script>`,全局名
   `pptxPreview`)。同时把这份 deck 拷进去。
2. 写一个 `index.html`:一个宽 360px 左右的 `div`,`pptxPreview.init(div,
   { width: 960, mode: "list" })` → `.preview(arrayBuffer)`,然后把
   `querySelectorAll` 数出来的东西写进页面里(几张片子、有没有横向溢出、文字在不在)。
   写一个静态服务器(`python -m http.server`)指向那个目录。
3. 浏览器打开它,**看截图** —— 空白帧 / 少一页 / 文字挤成一团,都只有看图才知道。

⚠️ 截图之前先给窗口一个真实尺寸。无头/隐藏标签页的视口是 0×0,那时候量出来的高度
全是 0,而"容器高度是 0"看起来和真 bug 一模一样。另外**浏览器面板会把整页缩着画**,
所以"看着小"不等于"渲染得小" —— 要判断尺寸,量 `getBoundingClientRect`,别用眼睛。

## 比肉眼更可靠的办法:拿 LibreOffice 当标准答案

"排得对不对"这种问题,和自己的直觉比不如和**同一个文件的另一种渲染**比。这台机器上
装了 LibreOffice(`C:\Program Files\LibreOffice\program\soffice.exe`),所以:

    soffice --headless --convert-to pdf -env:UserInstallation=file:///<一次性目录> \
            --outdir <目录> <一个 .pptx>
    # 再把 PDF 每页画成 PNG(pypdfium2 或 pdftoppm,这台机器上都有)

然后把自己库里排出来的那一页和它对一遍。**2026-09-16 就是这么发现 EMC.pptx 第 3 页
(那个 EMC/EMI/EMS 分组图)排得和 PowerPoint 不一样,而第 1/2/4 页基本一致** ——
纯前端库对分组图形的坐标处理有偏差,这不是我们这层能修的。

两个坑:LibreOffice **一个用户配置目录只允许跑一个实例**(并发转换会互相锁死,所以
`-env:UserInstallation` 指一个一次性的目录);第一次跑要建配置,冷启动十几秒。

## 两份稿子各压什么

`deck.pptx`(基准):标题页占位符、项目符号(含加粗 / 颜色)、表格、自选图形 + 文本框。
`--hard`:一张图(Pillow 生成的三色块,能一眼看出加载没加载)、一张图表(走 echarts)、
一页渐变背景上的白字。

产出的那两个 `.pptx` **不用提交** —— 它们是这个脚本的临时产物,和 `fixtures/` 里那三个
不是一回事:那三个由 `run.sh` 每次自动跑,所以进了仓库;这两份只在手工核对时用,跑一下
就有。
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
HARD = "--hard" in sys.argv
OUT = os.path.join(HERE, "deck-hard.pptx" if HARD else "deck.pptx")


def new_deck():
    prs = Presentation()
    # 16:9 —— 模版库里的演示稿基本都是这个比例,和默认的 4:3 排出来完全是两个样子。
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def base_deck():
    """四张片子,各压一种能力。"""
    prs = new_deck()

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "冒烟用演示文稿"
    s1.placeholders[1].text = "用来核对 pptx 到底排没排出来"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "这一页有什么"
    body = s2.placeholders[1].text_frame
    body.text = "第一行是普通的项目符号"
    p = body.add_paragraph()
    p.text = "第二行是加粗的"
    p.runs[0].font.bold = True
    p2 = body.add_paragraph()
    p2.text = "第三行是红色的"
    p2.runs[0].font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "一张表"
    tbl = s3.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(11), Inches(3)).table
    data = [["项目", "数值", "备注"], ["甲", "12.5", "上行"], ["乙", "47", "下行"]]
    for r in range(3):
        for c in range(3):
            tbl.cell(r, c).text = data[r][c]

    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    box.text_frame.text = "一个圆角矩形"
    tb = s4.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(1))
    tb.text_frame.text = "一个独立的文本框"
    return prs


def hard_deck():
    """再往上加三样模版里常见、但前一份没覆盖的东西。"""
    from PIL import Image, ImageDraw
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = new_deck()

    png = os.path.join(HERE, "_pic.png")
    img = Image.new("RGB", (600, 300), (240, 240, 240))
    d = ImageDraw.Draw(img)
    for i, c in enumerate([(200, 60, 60), (60, 160, 90), (60, 90, 200)]):
        d.rectangle([20 + i * 190, 40, 180 + i * 190, 260], fill=c)
    img.save(png)

    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "一张图"
    s1.shapes.add_picture(png, Inches(1), Inches(2), width=Inches(6))

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "一张图表"
    cd = CategoryChartData()
    cd.categories = ["甲", "乙", "丙"]
    cd.add_series("系列一", (12.5, 30, 47))
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2),
                        Inches(9), Inches(4), cd)

    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    sh = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    sh.fill.gradient()
    sh.fill.gradient_stops[0].color.rgb = RGBColor(0x1F, 0x3B, 0x73)
    sh.fill.gradient_stops[1].color.rgb = RGBColor(0x8E, 0x2B, 0x2B)
    sh.line.fill.background()
    tb = s3.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1.5))
    tb.text_frame.text = "渐变背景上的一行字"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    os.remove(png)
    return prs


prs = hard_deck() if HARD else base_deck()
prs.save(OUT)
print("wrote %s  %d bytes  %d slides  %.3f x %.3f in" % (
    OUT, os.path.getsize(OUT), len(prs.slides._sldIdLst),
    prs.slide_width / 914400, prs.slide_height / 914400))
